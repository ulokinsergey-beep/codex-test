#!/usr/bin/env python3
"""
Скрипт для сбора товаров с https://miuz.ru/catalog/.

Примеры запуска:
    python miuz_scraper.py
    python miuz_scraper.py --product-url https://miuz.ru/catalog/earrings/E01-EST-0246ES/

Нужные пакеты:
    pip install cloudscraper beautifulsoup4 openpyxl
"""

from __future__ import annotations

import argparse
import json
import random
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse
from xml.etree import ElementTree as ET

import cloudscraper
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook


BASE_URL = "https://miuz.ru"
CATALOG_URL = "https://miuz.ru/catalog/"
SITEMAP_URL = "https://miuz.ru/sitemap.xml"
OUTPUT_XLSX = "miuz_products.xlsx"
IMAGES_DIR = "images"
HEADERS = [
    "УИН",
    "ID",
    "Артикул",
    "Название изделия",
    "Цена",
    "Ссылка на товар",
]


class CaptchaError(RuntimeError):
    """Сайт вернул антибот-страницу вместо нужной страницы."""


@dataclass
class Product:
    uin: str
    product_id: str
    article: str
    name: str
    price: str
    url: str
    images: list[str]


class MiuzScraper:
    def __init__(
        self,
        output: Path,
        images_dir: Path,
        delay_min: float = 2.0,
        delay_max: float = 3.0,
        retries: int = 3,
        cookies: str | None = None,
    ) -> None:
        self.output = output
        self.images_dir = images_dir
        self.delay_min = delay_min
        self.delay_max = delay_max
        self.retries = retries
        self.last_request_at = 0.0
        self.session = cloudscraper.create_scraper(
            browser={
                "browser": "chrome",
                "platform": "windows",
                "desktop": True,
            }
        )
        self.session.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0.0.0 Safari/537.36"
                ),
                "Accept": (
                    "text/html,application/xhtml+xml,application/xml;q=0.9,"
                    "image/avif,image/webp,image/apng,*/*;q=0.8,"
                    "application/signed-exchange;v=b3;q=0.7"
                ),
                "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
                "Accept-Encoding": "gzip, deflate, br",
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "DNT": "1",
                "Pragma": "no-cache",
                "Priority": "u=0, i",
                "Referer": CATALOG_URL,
                "Sec-Ch-Ua": '"Chromium";v="124", "Google Chrome";v="124", "Not-A.Brand";v="99"',
                "Sec-Ch-Ua-Mobile": "?0",
                "Sec-Ch-Ua-Platform": '"Windows"',
                "Sec-Fetch-Dest": "document",
                "Sec-Fetch-Mode": "navigate",
                "Sec-Fetch-Site": "same-origin",
                "Sec-Fetch-User": "?1",
                "Upgrade-Insecure-Requests": "1",
            }
        )
        if cookies:
            self.session.cookies.update(parse_cookie_header(cookies))

    def run(
        self,
        product_url: str | None = None,
        limit: int | None = None,
        category_filter: str | None = None,
    ) -> None:
        workbook, sheet = self.open_workbook()
        processed_links = self.load_processed_links(sheet)

        if product_url:
            product_urls = [product_url]
        else:
            print("Собираю ссылки товаров из sitemap...")
            product_urls = self.discover_product_urls()
            total_found = len(product_urls)
            print(f"Всего товаров найдено: {total_found}")
            if category_filter:
                product_urls = [url for url in product_urls if category_filter in url]
                print(f"Товаров после фильтра {category_filter}: {len(product_urls)}")

        if limit is not None:
            product_urls = product_urls[:limit]

        total = len(product_urls)
        print(f"Найдено товаров для обхода: {total}")

        for index, url in enumerate(product_urls, start=1):
            normalized_url = normalize_url(url)
            if normalized_url in processed_links:
                print(f"Товар {index} из {total}: уже есть в Excel, пропускаю {normalized_url}")
                continue

            print(f"Товар {index} из {total}: загружаю {normalized_url}")
            try:
                product = self.parse_product(normalized_url)
                downloaded = self.download_images(product.images, product.article)
                self.append_product(sheet, product)
                workbook.save(self.output)
                processed_links.add(normalized_url)
                print(
                    "Готово: "
                    f"{product.article or 'без артикула'} | "
                    f"{product.name or 'без названия'} | "
                    f"фото скачано: {downloaded}"
                )
            except CaptchaError as exc:
                workbook.save(self.output)
                raise SystemExit(
                    "MIUZ вернул капчу вместо карточки товара. "
                    "Запустите скрипт с домашнего/офисного IP или позже. "
                    f"Подробность: {exc}"
                ) from exc
            except Exception as exc:  # noqa: BLE001 - парсер должен идти дальше по каталогу
                workbook.save(self.output)
                print(f"Ошибка на товаре {normalized_url}: {exc}", file=sys.stderr)

        workbook.save(self.output)
        print(f"Готово. Excel сохранён: {self.output}")

    def get(self, url: str, *, stream: bool = False) -> requests.Response:
        last_error: Exception | None = None

        for attempt in range(1, self.retries + 1):
            self.wait_before_request()
            try:
                response = self.session.get(url, timeout=40, stream=stream)
                self.last_request_at = time.monotonic()
                response.encoding = response.apparent_encoding or response.encoding

                if self.is_captcha(response):
                    raise CaptchaError(response.url)

                if response.status_code >= 500:
                    raise requests.HTTPError(f"HTTP {response.status_code}", response=response)

                response.raise_for_status()
                return response
            except CaptchaError:
                raise
            except Exception as exc:  # noqa: BLE001 - нужны повторы для сетевых ошибок
                last_error = exc
                print(f"Запрос не удался ({attempt}/{self.retries}): {url} — {exc}")

        raise RuntimeError(f"Страница не загрузилась после {self.retries} попыток: {url}") from last_error

    def wait_before_request(self) -> None:
        if not self.last_request_at:
            return

        pause = random.uniform(self.delay_min, self.delay_max)
        if pause > 0:
            time.sleep(pause)

    @staticmethod
    def is_captcha(response: requests.Response) -> bool:
        url = response.url.lower()
        content_type = response.headers.get("content-type", "").lower()
        text_start = ""
        if "text" in content_type or "html" in content_type or "xml" in content_type:
            text_start = response.text[:12000].lower()
        return (
            "showcaptcha" in url
            or "tmgrdfrend" in url
            or "smart-captcha" in text_start
            or "вы не робот" in text_start
        )

    def discover_product_urls(self) -> list[str]:
        sitemap_urls = self.parse_sitemap_index(SITEMAP_URL)
        product_urls: list[str] = []
        seen: set[str] = set()

        for sitemap_url in sitemap_urls:
            for loc in self.parse_sitemap_urls(sitemap_url):
                normalized = normalize_url(loc)
                if normalized in seen or not is_product_url(normalized):
                    continue
                seen.add(normalized)
                product_urls.append(normalized)

        return product_urls

    def parse_sitemap_index(self, sitemap_url: str) -> list[str]:
        response = self.get(sitemap_url)
        root = ET.fromstring(response.content)
        sitemap_urls = xml_values(root, "loc")

        if not sitemap_urls and sitemap_url.endswith(".xml"):
            return [sitemap_url]

        return sitemap_urls

    def parse_sitemap_urls(self, sitemap_url: str) -> list[str]:
        response = self.get(sitemap_url)
        root = ET.fromstring(response.content)

        if root.tag.endswith("sitemapindex"):
            urls: list[str] = []
            for nested_sitemap in xml_values(root, "loc"):
                urls.extend(self.parse_sitemap_urls(nested_sitemap))
            return urls

        return xml_values(root, "loc")

    def parse_product(self, url: str) -> Product:
        response = self.get(url)
        soup = BeautifulSoup(response.text, "html.parser")
        page_text = soup.get_text("\n", strip=True)
        json_objects = list(extract_json_objects(soup))
        product_json = first_product_json(json_objects)
        product_data = extract_nuxt_product_data(soup)
        selected_offer = find_selected_offer(product_data)

        article = (
            clean_value(product_data.get("code") or product_data.get("xmlId"))
            or clean_value(selected_offer.get("code") or selected_offer.get("xmlId"))
            or clean_value(selected_offer.get("offerXmlId"))
            or clean_value(value_from_json(product_json, "sku"))
            or find_labeled_value(page_text, ["Артикул", "SKU"])
            or article_from_url(url)
        )
        name = (
            clean_value(product_data.get("name"))
            or clean_value(product_data.get("seo", {}).get("h1") if isinstance(product_data.get("seo"), dict) else "")
            or clean_value(value_from_json(product_json, "name"))
            or text_of_first(soup, ["h1"])
            or clean_title(meta_content(soup, "og:title"))
        )
        price = (
            price_from_offer(selected_offer)
            or extract_price_from_json(product_json)
            or extract_visible_price(soup, page_text)
        )
        uin = (
            clean_value(selected_offer.get("uin"))
            or find_regex(response.text, [r"(?:УИН|UIN)[^\d]{0,40}(\d{8,})"])
            or find_labeled_value(page_text, ["УИН", "UIN"])
        )
        product_id = (
            clean_value(selected_offer.get("offerXmlId") or selected_offer.get("id"))
            or find_regex(
                response.text,
                [
                    r"(?:PRODUCT_ID|productId|product_id|data-product-id|itemId|offerId)[\"'\s:=,-]{1,20}(\d{4,})",
                    r"(?:ID товара|ID изделия)[^\d]{0,40}(\d{4,})",
                ],
            )
            or find_labeled_value(page_text, ["ID товара", "ID изделия", "ID"])
        )
        images = extract_images(soup, response.text, json_objects, url)

        missing = []
        for label, value in [
            ("УИН", uin),
            ("ID", product_id),
            ("Артикул", article),
            ("Название", name),
            ("Цена", price),
        ]:
            if not value:
                missing.append(label)
        if missing:
            print(f"Предупреждение: не найдены поля: {', '.join(missing)}")

        return Product(
            uin=uin,
            product_id=product_id,
            article=article,
            name=name,
            price=price,
            url=normalize_url(url),
            images=images,
        )

    def download_images(self, image_urls: Iterable[str], article: str) -> int:
        safe_article = safe_path_name(article or "without_article")
        target_dir = self.images_dir / safe_article
        target_dir.mkdir(parents=True, exist_ok=True)

        downloaded = 0
        for index, image_url in enumerate(unique(image_urls), start=1):
            try:
                ext = image_extension(image_url)
                path = target_dir / f"{index:02d}{ext}"
                if path.exists() and path.stat().st_size > 0:
                    downloaded += 1
                    continue

                response = self.get(image_url, stream=True)
                content_type = response.headers.get("content-type", "")
                if "image" not in content_type.lower():
                    print(f"Пропускаю не-картинку: {image_url}")
                    continue

                with path.open("wb") as file:
                    for chunk in response.iter_content(chunk_size=65536):
                        if chunk:
                            file.write(chunk)
                downloaded += 1
            except Exception as exc:  # noqa: BLE001 - одна битая картинка не должна ломать весь сбор
                print(f"Не удалось скачать фото {image_url}: {exc}", file=sys.stderr)

        return downloaded

    def open_workbook(self):
        if self.output.exists():
            workbook = load_workbook(self.output)
            sheet = workbook.active
            if sheet.max_row < 1:
                sheet.append(HEADERS)
            return workbook, sheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MIUZ products"
        sheet.append(HEADERS)
        workbook.save(self.output)
        return workbook, sheet

    @staticmethod
    def load_processed_links(sheet) -> set[str]:
        header = [cell.value for cell in sheet[1]]
        try:
            link_column = header.index("Ссылка на товар") + 1
        except ValueError:
            return set()

        links: set[str] = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[link_column - 1]
            if value:
                links.add(normalize_url(str(value)))
        return links

    @staticmethod
    def append_product(sheet, product: Product) -> None:
        sheet.append(
            [
                product.uin,
                product.product_id,
                product.article,
                product.name,
                product.price,
                product.url,
            ]
        )


def xml_values(root: ET.Element, tag_name: str) -> list[str]:
    values: list[str] = []
    for node in root.iter():
        if node.tag.endswith(tag_name) and node.text:
            values.append(node.text.strip())
    return values


def is_product_url(url: str) -> bool:
    parsed = urlparse(url)
    parts = [part for part in parsed.path.strip("/").split("/") if part]
    if len(parts) < 3 or parts[0] != "catalog":
        return False
    if "filter" in parts:
        return False

    article = parts[-1]
    if article in {"new", "sale", "hit", "diamonds"}:
        return False

    # У товаров MIUZ в ссылке обычно артикул: буквы/цифры/дефисы и хотя бы одна цифра.
    return bool(re.search(r"\d", article)) and bool(re.fullmatch(r"[A-Za-zА-Яа-я0-9_-]+", article))


def normalize_url(url: str) -> str:
    url = urljoin(BASE_URL, url)
    parsed = urlparse(url)
    path = parsed.path
    if not path.endswith("/"):
        path += "/"
    return f"{parsed.scheme}://{parsed.netloc}{path}"


def article_from_url(url: str) -> str:
    return urlparse(url).path.rstrip("/").split("/")[-1]


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        value = next((item for item in value if item), "")
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_title(value: str) -> str:
    value = clean_value(value)
    value = re.sub(r"\s+по цене от\s+.+$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\s+-\s+купить.+$", "", value, flags=re.IGNORECASE)
    return value.strip()


def text_of_first(soup: BeautifulSoup, selectors: Iterable[str]) -> str:
    for selector in selectors:
        node = soup.select_one(selector)
        if node:
            text = clean_value(node.get_text(" ", strip=True))
            if text:
                return text
    return ""


def meta_content(soup: BeautifulSoup, property_name: str) -> str:
    node = soup.find("meta", attrs={"property": property_name}) or soup.find(
        "meta", attrs={"name": property_name}
    )
    return clean_value(node.get("content")) if node else ""


def find_labeled_value(text: str, labels: Iterable[str]) -> str:
    for label in labels:
        pattern = rf"{re.escape(label)}\s*[:№#]?\s*([\wА-Яа-яЁё./-]+)"
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return clean_value(match.group(1))
    return ""


def find_regex(text: str, patterns: Iterable[str]) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if match:
            return clean_value(match.group(1))
    return ""


def extract_json_objects(soup: BeautifulSoup) -> Iterable[Any]:
    for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
        raw = script.string or script.get_text(strip=True)
        if not raw:
            continue
        try:
            yield json.loads(raw)
        except json.JSONDecodeError:
            continue

    for script in soup.find_all("script"):
        raw = script.string or script.get_text()
        if not raw:
            continue

        for pattern in [
            r"window\.__INITIAL_STATE__\s*=\s*({.*?})\s*;",
            r"window\.__NUXT__\s*=\s*({.*?})\s*;",
            r"window\.__NEXT_DATA__\s*=\s*({.*?})\s*;",
        ]:
            match = re.search(pattern, raw, flags=re.DOTALL)
            if not match:
                continue
            try:
                yield json.loads(match.group(1))
            except json.JSONDecodeError:
                continue


def iter_json_values(data: Any) -> Iterable[Any]:
    if isinstance(data, dict):
        yield data
        for value in data.values():
            yield from iter_json_values(value)
    elif isinstance(data, list):
        for item in data:
            yield from iter_json_values(item)


def first_product_json(json_objects: Iterable[Any]) -> dict[str, Any]:
    for obj in json_objects:
        for value in iter_json_values(obj):
            if not isinstance(value, dict):
                continue
            item_type = value.get("@type") or value.get("type")
            if isinstance(item_type, list):
                item_type = " ".join(map(str, item_type))
            if item_type and "product" in str(item_type).lower():
                return value
    return {}


def value_from_json(data: dict[str, Any], key: str) -> Any:
    if not data:
        return ""
    value = data.get(key)
    if value:
        return value
    for nested in iter_json_values(data):
        if isinstance(nested, dict) and nested.get(key):
            return nested.get(key)
    return ""


def extract_price_from_json(product_json: dict[str, Any]) -> str:
    if not product_json:
        return ""

    offers = product_json.get("offers")
    offer_items = offers if isinstance(offers, list) else [offers]
    for offer in offer_items:
        if not isinstance(offer, dict):
            continue
        price = offer.get("price") or offer.get("lowPrice")
        if price:
            return format_price(price)
    return ""


def extract_visible_price(soup: BeautifulSoup, page_text: str) -> str:
    for selector in [
        '[itemprop="price"]',
        '[data-price]',
        ".price",
        ".product-price",
        ".catalog-element-price",
    ]:
        node = soup.select_one(selector)
        if not node:
            continue
        value = node.get("content") or node.get("data-price") or node.get_text(" ", strip=True)
        value = clean_value(value)
        if value:
            return format_price(value)

    match = re.search(r"(\d[\d\s\u00a0]{2,}\s*₽)", page_text)
    return clean_value(match.group(1)) if match else ""


def format_price(value: Any) -> str:
    text = clean_value(value).replace("\u00a0", " ")
    if "₽" in text or "руб" in text.lower():
        return text

    digits = re.sub(r"[^\d]", "", text)
    if not digits:
        return text
    return f"{int(digits):,}".replace(",", " ") + " ₽"


def extract_images(
    soup: BeautifulSoup,
    html: str,
    json_objects: Iterable[Any],
    page_url: str,
) -> list[str]:
    urls: list[str] = []

    # Nuxt keeps the current product images in /api/product/info. Prefer it so
    # related products, banners, and responsive duplicates are not downloaded.
    product_data = extract_nuxt_product_data(soup)
    urls.extend(extract_product_image_urls(product_data))
    if urls:
        return normalize_image_urls(urls, page_url)

    for obj in json_objects:
        for value in iter_json_values(obj):
            if not isinstance(value, dict):
                continue
            image_value = value.get("image") or value.get("images") or value.get("picture")
            urls.extend(as_list(image_value))

    for property_name in ["og:image", "og:image:secure_url"]:
        content = meta_content(soup, property_name)
        if content:
            urls.append(content)

    for node in soup.find_all(["img", "source"]):
        for attr in ["src", "data-src", "data-original", "data-lazy", "srcset", "data-srcset"]:
            value = node.get(attr)
            if not value:
                continue
            urls.extend(split_srcset(value))

    urls.extend(re.findall(r"https?://[^\s\"']+\.(?:jpg|jpeg|png|webp)(?:\?[^\s\"']*)?", html, re.I))
    urls.extend(re.findall(r"//[^\s\"']+\.(?:jpg|jpeg|png|webp)(?:\?[^\s\"']*)?", html, re.I))

    return normalize_image_urls(urls, page_url)


def normalize_image_urls(urls: Iterable[str], page_url: str) -> list[str]:
    cleaned: list[str] = []
    for url in urls:
        absolute = normalize_image_url(urljoin(page_url, url.strip()))
        if is_product_image_url(absolute):
            cleaned.append(absolute)

    return unique(cleaned)


def extract_nuxt_product_data(soup: BeautifulSoup) -> dict[str, Any]:
    for script in soup.find_all("script", attrs={"type": "application/json"}):
        raw = script.string or script.get_text(strip=True)
        if not raw:
            continue

        try:
            payload = decode_nuxt_payload(raw)
        except (json.JSONDecodeError, IndexError, TypeError, RecursionError):
            continue

        data = payload.get("data") if isinstance(payload, dict) else None
        if not isinstance(data, dict):
            continue

        for key, value in data.items():
            if "/api/product/info" not in key or not isinstance(value, dict):
                continue
            product_data = value.get("data")
            if isinstance(product_data, dict):
                return product_data

    return {}


def decode_nuxt_payload(raw: str) -> Any:
    values = json.loads(raw)
    if not isinstance(values, list):
        return values

    memo: dict[int, Any] = {}
    wrappers = {"Reactive", "ShallowReactive", "Readonly", "ShallowReadonly", "Ref", "ComputedRef"}

    def revive_ref(index: int) -> Any:
        if index in memo:
            return memo[index]
        value = values[index]
        memo[index] = None
        result = revive_value(value)
        memo[index] = result
        return result

    def revive_child(value: Any) -> Any:
        if type(value) is int and 0 <= value < len(values):
            return revive_ref(value)
        if type(value) is int:
            return value
        return revive_value(value)

    def revive_value(value: Any) -> Any:
        if type(value) is int:
            return value
        if isinstance(value, list):
            if value and value[0] in wrappers and len(value) > 1:
                return revive_child(value[1])
            return [revive_child(item) for item in value]
        if isinstance(value, dict):
            return {key: revive_child(item) for key, item in value.items()}
        return value

    return revive_ref(0)


def extract_product_image_urls(product_data: dict[str, Any]) -> list[str]:
    image_urls: list[str] = []
    images = product_data.get("images")
    if not isinstance(images, list):
        return image_urls

    for image in images:
        if not isinstance(image, dict) or image.get("duplicate"):
            continue
        selected_url = ""
        for group_name in ["main", "gallery", "original", "thumb", "galleryMin"]:
            group = image.get(group_name)
            if not isinstance(group, dict):
                continue
            for key in ["src", "srcX2"]:
                value = group.get(key)
                if value:
                    selected_url = str(value)
                    break
            if selected_url:
                break
        if selected_url:
            image_urls.append(selected_url)

    return image_urls


def find_selected_offer(product_data: dict[str, Any]) -> dict[str, Any]:
    for value in iter_json_values(product_data):
        if isinstance(value, dict) and value.get("selected") is True and value.get("offerXmlId"):
            return value
    return {}


def price_from_offer(offer: dict[str, Any]) -> str:
    price = offer.get("price")
    if isinstance(price, dict):
        value = price.get("value")
        currency_sign = price.get("currencySign") or "₽"
        if value:
            return f"{int(value):,}".replace(",", " ") + f" {currency_sign}"
    if price:
        return format_price(price)
    return ""


def as_list(value: Any) -> list[str]:
    if not value:
        return []
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [clean_value(value.get("url") or value.get("contentUrl"))]
    if isinstance(value, list):
        items: list[str] = []
        for item in value:
            items.extend(as_list(item))
        return items
    return []


def split_srcset(value: str) -> list[str]:
    result: list[str] = []
    for part in value.split(","):
        url = part.strip().split(" ")[0]
        if url:
            result.append(url)
    return result


def normalize_image_url(url: str) -> str:
    if url.startswith("//"):
        url = "https:" + url
    parsed = urlparse(url)
    return f"{parsed.scheme}://{parsed.netloc}{parsed.path}" + (f"?{parsed.query}" if parsed.query else "")


def is_product_image_url(url: str) -> bool:
    lower = url.lower()
    if not re.search(r"\.(jpg|jpeg|png|webp)(?:\?|$)", lower):
        return False
    blocked = ["logo", "icon", "sprite", "captcha", "placeholder", "loader", "favicon"]
    if any(word in lower for word in blocked):
        return False
    return "/upload/" in lower or "miuz" in lower or "cdn" in lower


def image_extension(url: str) -> str:
    path = urlparse(url).path.lower()
    match = re.search(r"\.(jpg|jpeg|png|webp)$", path)
    if not match:
        return ".jpg"
    ext = match.group(1)
    return ".jpg" if ext == "jpeg" else f".{ext}"


def safe_path_name(value: str) -> str:
    value = clean_value(value)
    value = re.sub(r'[<>:"/\\|?*]+', "_", value)
    value = value.strip(" .")
    return value or "without_article"


def unique(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = clean_value(value)
        if not value or value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def parse_cookie_header(cookie_header: str) -> dict[str, str]:
    cookie_header = cookie_header.strip()
    if cookie_header.lower().startswith("cookie:"):
        cookie_header = cookie_header.split(":", 1)[1].strip()

    cookies: dict[str, str] = {}
    for part in cookie_header.split(";"):
        if "=" not in part:
            continue
        name, value = part.split("=", 1)
        name = name.strip()
        if not name:
            continue
        cookies[name] = value.strip()
    return cookies


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Парсер товаров MIUZ")
    parser.add_argument("--product-url", help="Собрать только один товар по ссылке")
    parser.add_argument("--limit", type=int, help="Ограничить количество товаров для проверки")
    parser.add_argument(
        "--category-filter",
        help='Фильтр категории по фрагменту URL, например "/catalog/earrings/"',
    )
    parser.add_argument("--output", default=OUTPUT_XLSX, help="Файл Excel")
    parser.add_argument("--images-dir", default=IMAGES_DIR, help="Папка для фото")
    parser.add_argument("--delay-min", type=float, default=2.0, help="Минимальная задержка между запросами")
    parser.add_argument("--delay-max", type=float, default=3.0, help="Максимальная задержка между запросами")
    parser.add_argument("--retries", type=int, default=3, help="Количество повторов запроса")
    parser.add_argument(
        "--cookies",
        "--cookie",
        dest="cookies",
        help='Cookies из браузера, например: "name=value; name2=value2"',
    )
    return parser.parse_args()


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(line_buffering=True)
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(line_buffering=True)

    args = parse_args()
    if args.delay_min < 0 or args.delay_max < 0 or args.delay_max < args.delay_min:
        raise SystemExit("Некорректные задержки: нужно 0 <= delay-min <= delay-max")

    scraper = MiuzScraper(
        output=Path(args.output),
        images_dir=Path(args.images_dir),
        delay_min=args.delay_min,
        delay_max=args.delay_max,
        retries=args.retries,
        cookies=args.cookies,
    )
    scraper.run(product_url=args.product_url, limit=args.limit, category_filter=args.category_filter)


if __name__ == "__main__":
    main()
